# -*- coding: utf-8 -*-
"""
Atualizador Dashboard Firjan - IMES
Cards e graficos vem do arquivo de Status (aba Resumo: tabelas Controle + Lista de Espera
somadas). Tentativas e periodo vem da discagem.
"""

import openpyxl
import os
import re
import glob
import unicodedata
from datetime import datetime

PASTA        = r'Arquivos\nao_atualizaveis\Ativo_IMES'
PREFIXO_DISC = 'DISCAGEM_IMES'
PREFIXO_BASE = 'Status - '
INDEX_HTML   = r'index.html'
DEPARA_PATH  = r'Arquivos\bases_apoio\tab_de-para.xlsx'

COL_DATA = 0

# Card "Interessados" = somente estes status de sucesso
INTERESSADOS = ['Confirmado']

GRUPO_SUCESSO   = 'sucesso no contato'
GRUPO_INSUCESSO = 'insucesso no contato'


def norm_txt(s):
    if s is None:
        return ''
    s = str(s).strip().upper()
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def encontrar_arquivo(pasta, prefixo):
    padrao = os.path.join(pasta, f'{prefixo}*.xlsx')
    arquivos = [a for a in glob.glob(padrao) if not os.path.basename(a).startswith('~$')]
    return sorted(arquivos)[-1] if arquivos else None


def data_int(dt):
    if isinstance(dt, datetime):
        return dt.year * 10000 + dt.month * 100 + dt.day
    return None


def ler_depara(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    mapa = {}
    for r in wb['status'].iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[1]:
            mapa[norm_txt(r[0])] = str(r[1]).strip()
    wb.close()
    return mapa


def ler_resumo(caminho):
    """Soma as tabelas Controle e Lista de Espera da aba Resumo.
    Retorna (sucesso{nome:qtd}, insucesso{nome:qtd}, motivos{nome:qtd}, total_geral)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    linhas = [list(r) for r in wb['Resumo'].iter_rows(values_only=True)]
    wb.close()

    sucesso, insucesso, motivos = {}, {}, {}
    nomes = {}
    total_geral = 0

    def acumula(dic, nome, qtd):
        k = norm_txt(nome)
        nomes.setdefault(k, str(nome).strip())
        dic[k] = dic.get(k, 0) + qtd

    # Cada tabela ocupa um par de colunas (rotulo, quantidade). Descobre pares pelo cabecalho.
    for i, row in enumerate(linhas):
        for c, cel in enumerate(row):
            t = norm_txt(cel)
            if t not in ('CONTATO', 'MOTIVO DO NAO INTERESSE'):
                continue
            eh_contato = (t == 'CONTATO')
            grupo = None
            for r2 in linhas[i + 1:]:
                nome = r2[c] if len(r2) > c else None
                qtd = r2[c + 1] if len(r2) > c + 1 else None
                if nome is None:
                    break
                k = norm_txt(nome)
                if k == 'TOTAL GERAL':
                    if eh_contato and isinstance(qtd, (int, float)):
                        total_geral += int(qtd)
                    break
                if not isinstance(qtd, (int, float)):
                    continue
                qtd = int(qtd)
                if eh_contato:
                    if k == norm_txt(GRUPO_SUCESSO):
                        grupo = 'S'
                        continue
                    if k == norm_txt(GRUPO_INSUCESSO):
                        grupo = 'I'
                        continue
                    acumula(sucesso if grupo == 'S' else insucesso, nome, qtd)
                else:
                    acumula(motivos, nome, qtd)

    def rotula(d):
        return {nomes[k]: v for k, v in d.items()}

    return rotula(sucesso), rotula(insucesso), rotula(motivos), total_geral


def calcular_imes(caminho_disc, caminho_base):
    print(f'  Status: {os.path.basename(caminho_base)}')
    sucesso, insucesso, motivos, total_emp = ler_resumo(caminho_base)
    print(f'  Empresas (Total Geral Controle + Lista de Espera): {total_emp}')
    print(f'  Sucesso: {sucesso}')
    print(f'  Insucesso: {insucesso}')
    print(f'  Motivos: {motivos}')

    try:
        depara = ler_depara(DEPARA_PATH)
        print(f'  De-para: {len(depara)} entradas')
    except Exception as e:
        depara = {}
        print(f'  [AVISO] De-para nao carregado: {e}')

    # Insucesso agrupado pelo de-para
    ns, raw = {}, {}
    for nome, qtd in insucesso.items():
        lbl = depara.get(norm_txt(nome), nome)
        ns[lbl] = ns.get(lbl, 0) + qtd
        raw.setdefault(lbl, []).append(nome)
    ns_ord = sorted(ns.items(), key=lambda kv: -kv[1])

    suc_ord = sorted(sucesso.items(), key=lambda kv: -kv[1])
    mot_ord = sorted(motivos.items(), key=lambda kv: -kv[1])

    # Tentativas e periodo (discagem)
    tentativas, datas = 0, []
    if caminho_disc:
        print(f'  Discagem: {os.path.basename(caminho_disc)}')
        wb = openpyxl.load_workbook(caminho_disc, read_only=True, data_only=True)
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            dt = data_int(row[COL_DATA])
            if dt is None:
                continue
            datas.append(dt)
            tentativas += 1
        wb.close()
    periodo = ''
    if datas:
        a, b = min(datas), max(datas)
        periodo = f"{a % 100:02d}/{(a // 100) % 100:02d}/{a // 10000} — {b % 100:02d}/{(b // 100) % 100:02d}/{b // 10000}"
    print(f'  Tentativas: {tentativas} | Periodo: {periodo}')

    int_norm = {norm_txt(x) for x in INTERESSADOS}
    return {
        'empresas': total_emp,
        'tentativas': tentativas,
        'periodo': periodo,
        'sucessoLabels': [k for k, _ in suc_ord],
        'sucessoData': [v for _, v in suc_ord],
        'decisor': sum(v for _, v in suc_ord),
        'interessados': sum(v for k, v in suc_ord if norm_txt(k) in int_norm),
        'naosucessoLabels': [k for k, _ in ns_ord],
        'naosucessoData': [v for _, v in ns_ord],
        'naosucessoTooltips': [', '.join(raw[k]) for k, _ in ns_ord],
        'motivoLabels': [k for k, _ in mot_ord],
        'motivoData': [v for _, v in mot_ord],
    }


def gerar_bloco(d):
    def js_str(lst):
        return '[' + ','.join(f"'{str(v).replace(chr(39), chr(92)+chr(39))}'" for v in lst) + ']'

    def js_num(lst):
        return '[' + ','.join(str(v) for v in lst) + ']'

    media = f"{d['tentativas'] / d['empresas']:.2f}".replace('.', ',') if d['empresas'] else '—'
    conv = (f"{d['interessados'] / d['decisor'] * 100:.2f}".replace('.', ',') + '%') if d['decisor'] else '—'
    emp = str(d['empresas']) if d['empresas'] else '—'
    mostra_motivo = 'true' if d['motivoLabels'] else 'false'

    return f"""  /* IMES_START */
  imes: {{
    label: '— IMES', desc: 'Campanha IMES — dados filtrados', periodo: '{d['periodo']}',
    empresas: '{emp}', empresasLabel: '🏢 Empresas',
    mediaLabel: '🔁 Média', mediaSub: 'por empresa',
    tentativas: '{d['tentativas']}', interessados: '{d['interessados']}', conversao: '{conv}',
    decisor: '{d['decisor']}', decisorLabel: '👤 Contatos de Sucesso', decisorSub: 'IMES', media: '{media}', trend: '',
    distTitle: 'Contatos de Sucesso',
    statusLabels: {js_str(d['sucessoLabels'])}, statusData: {js_num(d['sucessoData'])}, statusColors: null,
    statusTooltips: {js_str(d['sucessoLabels'])},
    evoTitle: 'Tentativas de Contato Sem Sucesso', evoBar: true,
    naosucessoLabels: {js_str(d['naosucessoLabels'])}, naosucessoData: {js_num(d['naosucessoData'])},
    naosucessoTooltips: {js_str(d['naosucessoTooltips'])},
    evolucaoLabels: [], tentDia: [], convDia: [],
    showWpp: false,
    wppTitle: '', wppDesc: '',
    wppKpiLabels: [], wppListLabels: [], wppPieLabels: [],
    wppEnv:'-', wppResp:'-', wppTaxa:'-', wppSem:'-', wppInfo:'-', wppEmail:'-', wppPie:[0,1],
    distToggle: false,
    showMotivo: {mostra_motivo},
    motivoLabels: {js_str(d['motivoLabels'])}, motivoData: {js_num(d['motivoData'])}
  }},
  /* IMES_END */"""


def atualizar_html(index_path, bloco):
    with open(index_path, 'r', encoding='utf-8') as f:
        conteudo = f.read()
    padrao = r'/\* IMES_START \*/.*?/\* IMES_END \*/'
    if re.search(padrao, conteudo, re.DOTALL):
        conteudo = re.sub(padrao, lambda m: bloco, conteudo, flags=re.DOTALL)
    elif '/* POTENCIALIZEE_START */' in conteudo:
        conteudo = conteudo.replace('/* POTENCIALIZEE_START */', bloco + '\n  /* POTENCIALIZEE_START */', 1)
    else:
        raise ValueError('[ERRO] Marcadores nao encontrados no index.html.')
    with open(index_path, 'w', encoding='utf-8') as f:
        f.write(conteudo)


def main():
    print()
    print('=' * 50)
    print('  ATUALIZADOR — IMES')
    print('=' * 50)
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    try:
        caminho_base = encontrar_arquivo(PASTA, PREFIXO_BASE)
        if not caminho_base:
            print(f'  [AVISO] Arquivo de status nao encontrado em {PASTA}')
            return
        caminho_disc = encontrar_arquivo(PASTA, PREFIXO_DISC)
        dados = calcular_imes(caminho_disc, caminho_base)
        atualizar_html(INDEX_HTML, gerar_bloco(dados))
        print()
        print('=' * 50)
        print('  CONCLUIDO! index.html atualizado.')
        print('  Rode publicar.bat para enviar ao GitHub.')
        print('=' * 50)
        print()
    except Exception as e:
        print(f'\n[ERRO] {e}')
        import traceback
        traceback.print_exc()


if __name__ == '__main__':
    main()
