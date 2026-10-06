# -*- coding: utf-8 -*-
"""
Atualizador Dashboard Firjan - Propag Nova Iguaçu
Le discagem e base, atualiza o bloco PROPAG_NI no index.html.

Logica especial: Status "Atendido" pega STATUS_NEGOCIO (ou fica "Atendido" se vazio)
"""

import openpyxl
import os
import re
import glob
from datetime import datetime

PASTA      = r'Arquivos\nao_atualizaveis\Ativo_PROPAG NI'
PREFIXO_DISC = 'Discagem_propag_NOVA'
PREFIXO_BASE = 'Status - '
INDEX_HTML = r'index.html'

# Colunas (0-indexed)
COL_DATA = 0
COL_STATUS = 10
COL_STATUS_NEGOCIO = 11


# Status (ou STATUS_NEGOCIO, quando Atendido) que contam como "Contatos de Sucesso".
# Escreva exatamente como aparece na discagem (maiusculas/acentos sao ignorados).
SUCESSO_LABELS = []

SEM_OPERADOR_LABELS = ['Falhou', 'Fora de Area / Cx de Mensagens', 'Ligação Muda', 'Não Atendeu',
                       'Ocupado', 'Tel Não Atende / Ocupado', 'Atendido']


def encontrar_arquivo(pasta, prefixo):
    padrao = os.path.join(pasta, f'{prefixo}*.xlsx')
    arquivos = [a for a in glob.glob(padrao) if not os.path.basename(a).startswith('~$')]
    if not arquivos:
        return None
    return sorted(arquivos)[-1]


def data_int(dt):
    if dt is None:
        return None
    if isinstance(dt, datetime):
        return dt.year * 10000 + dt.month * 100 + dt.day
    s = str(dt).strip()
    if not s or s == '--':
        return None
    # Tenta "DD/MM/YYYY"
    try:
        parts = s.split('/')
        if len(parts) == 3:
            dd, mo, yy = int(parts[0]), int(parts[1]), int(parts[2])
            return yy * 10000 + mo * 100 + dd
    except:
        pass
    return None


def norm_txt(s):
    """Normaliza para uppercase sem acentos."""
    if not s:
        return ''
    import unicodedata
    s = str(s).strip().upper()
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


SEM_OPERADOR_NORM = {norm_txt(s) for s in SEM_OPERADOR_LABELS}
SUCESSO_MAP = {norm_txt(s): s for s in SUCESSO_LABELS}


def calcular_propag_ni(caminho_disc, caminho_base):
    print(f'  Discagem: {os.path.basename(caminho_disc)}')
    
    # Contar empresas na base
    total_emp = 0
    if caminho_base:
        print(f'  Base: {os.path.basename(caminho_base)}')
        try:
            wb_base = openpyxl.load_workbook(caminho_base, read_only=True, data_only=True)
            ws_base = wb_base['Base']
            linhas_base = [r for r in ws_base.iter_rows(values_only=True) if r and any(r)]
            hdr = [str(h).strip().upper() if h else '' for h in linhas_base[0]]
            i_insc = hdr.index('INSCRICAO') if 'INSCRICAO' in hdr else None
            if i_insc is None:
                total_emp = len(linhas_base) - 1
            else:
                total_emp = len({r[i_insc] for r in linhas_base[1:] if r[i_insc] is not None})
            wb_base.close()
            print(f'  Empresas na base: {total_emp}')
        except Exception as e:
            print(f'  [AVISO] Base não processada: {e}')
    
    wb = openpyxl.load_workbook(caminho_disc, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))[1:]

    total_tent = 0
    naosucesso_counter = {}
    sucesso_counter = {}
    raw_por_label = {}
    periodo_datas = []

    for row in rows:
        if not row or len(row) < 12:
            continue
        dt = data_int(row[COL_DATA])
        if dt is None:
            continue
        periodo_datas.append(dt)

        status_raw = str(row[COL_STATUS]).strip() if row[COL_STATUS] else ''
        status_negocio = str(row[COL_STATUS_NEGOCIO]).strip() if len(row) > COL_STATUS_NEGOCIO and row[COL_STATUS_NEGOCIO] else ''

        # Se Atendido, pega STATUS_NEGOCIO ou deixa "Atendido"
        if norm_txt(status_raw) == norm_txt('Atendido'):
            status_display = status_negocio if status_negocio else 'Atendido'
        else:
            status_display = status_raw

        if norm_txt(status_display) in SUCESSO_MAP:
            lbl = SUCESSO_MAP[norm_txt(status_display)]
            sucesso_counter[lbl] = sucesso_counter.get(lbl, 0) + 1
            total_tent += 1
            continue
        if norm_txt(status_display) in SEM_OPERADOR_NORM:
            label_ns = 'Tentativa'
            raw_por_label.setdefault(label_ns, [])
            if status_display not in raw_por_label[label_ns]:
                raw_por_label[label_ns].append(status_display)
        else:
            label_ns = status_display
        naosucesso_counter[label_ns] = naosucesso_counter.get(label_ns, 0) + 1
        total_tent += 1

    wb.close()

    periodo = ''
    if periodo_datas:
        dmin, dmax = min(periodo_datas), max(periodo_datas)
        d_min_str = f"{dmin % 100:02d}/{(dmin // 100) % 100:02d}/{dmin // 10000}"
        d_max_str = f"{dmax % 100:02d}/{(dmax // 100) % 100:02d}/{dmax // 10000}"
        periodo = f"{d_min_str} — {d_max_str}"

    print(f'  Total tentativas: {total_tent}')
    suc = sorted(sucesso_counter.items(), key=lambda kv: -kv[1])
    ordenado = sorted(naosucesso_counter.items(), key=lambda kv: -kv[1])
    ns_labels = [k for k, _ in ordenado]
    ns_data = [v for _, v in ordenado]
    ns_tooltips = [', '.join(raw_por_label.get(k, [k])) for k in ns_labels]
    print(f'  Sem sucesso: {dict(ordenado)}')
    print(f'  Periodo: {periodo}')

    return {
        'empresas': total_emp,
        'tentativas': total_tent,
        'sucessoLabels': [k for k, _ in suc],
        'sucessoData': [v for _, v in suc],
        'decisor': sum(v for _, v in suc),
        'naosucessoLabels': ns_labels,
        'naosucessoData': ns_data,
        'naosucessoTooltips': ns_tooltips,
        'periodo': periodo,
    }


def gerar_bloco(dados):
    def js_str(lst):
        return '[' + ','.join(f"'{str(v).replace(chr(39), chr(92)+chr(39))}'" for v in lst) + ']'

    def js_num(lst):
        return '[' + ','.join(str(v) for v in lst) + ']'

    media = (f"{dados['tentativas'] / dados['empresas']:.2f}".replace('.', ',')) if dados['empresas'] else '—'
    emp_display = str(dados['empresas']) if dados['empresas'] > 0 else '—'
    
    return f"""  /* PROPAG_NI_START */
  propag_ni: {{
    label: '— Propag Nova Iguaçu', desc: 'Campanha Propag Nova Iguaçu — dados filtrados', periodo: '{dados['periodo']}',
    empresas: '{emp_display}', empresasLabel: '🏢 Empresas',
    mediaLabel: '🔁 Média', mediaSub: 'por empresa',
    tentativas: '{dados['tentativas']}', interessados: '—', conversao: '—',
    decisor: '{dados['decisor']}', decisorLabel: '👤 Contatos de Sucesso', decisorSub: 'Propag Nova Iguaçu', media: '{media}', trend: '',
    distTitle: 'Contatos de Sucesso',
    statusLabels: {js_str(dados['sucessoLabels'])}, statusData: {js_num(dados['sucessoData'])}, statusColors: null,
    statusTooltips: {js_str(dados['sucessoLabels'])},
    evoTitle: 'Tentativas de Contato Sem Sucesso', evoBar: true,
    naosucessoLabels: {js_str(dados['naosucessoLabels'])}, naosucessoData: {js_num(dados['naosucessoData'])},
    naosucessoTooltips: {js_str(dados['naosucessoTooltips'])},
    evolucaoLabels: [], tentDia: [], convDia: [],
    showWpp: false,
    wppTitle: '', wppDesc: '',
    wppKpiLabels: [], wppListLabels: [], wppPieLabels: [],
    wppEnv:'-', wppResp:'-', wppTaxa:'-', wppSem:'-', wppInfo:'-', wppEmail:'-', wppPie:[0,1],
    distToggle: false,
    showMotivo: false,
    motivoLabels: [], motivoData: []
  }},
  /* PROPAG_NI_END */"""


def atualizar_html(index_path, bloco):
    with open(index_path, 'r', encoding='utf-8') as f:
        conteudo = f.read()
    padrao = r'/\* PROPAG_NI_START \*/.*?/\* PROPAG_NI_END \*/'
    if not re.search(padrao, conteudo, re.DOTALL):
        print('  [AVISO] Marcadores PROPAG_NI não encontrados, criando...')
        # Insere antes de POTENCIALIZEE_START se existir
        if '/* POTENCIALIZEE_START */' in conteudo:
            conteudo = conteudo.replace('/* POTENCIALIZEE_START */', bloco + '\n  /* POTENCIALIZEE_START */')
        else:
            raise ValueError('[ERRO] Nem PROPAG_NI nem POTENCIALIZEE marcadores encontrados.')
    else:
        conteudo = re.sub(padrao, lambda m: bloco, conteudo, flags=re.DOTALL)
    with open(index_path, 'w', encoding='utf-8') as f:
        f.write(conteudo)


def main():
    print()
    print('=' * 50)
    print('  ATUALIZADOR — Propag Nova Iguaçu')
    print('=' * 50)
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    try:
        caminho_disc = encontrar_arquivo(PASTA, PREFIXO_DISC)
        if not caminho_disc:
            print(f'  [AVISO] Discagem não encontrada em {PASTA}')
            return
        
        caminho_base = encontrar_arquivo(PASTA, PREFIXO_BASE)

        dados = calcular_propag_ni(caminho_disc, caminho_base)
        bloco = gerar_bloco(dados)
        atualizar_html(INDEX_HTML, bloco)
        print()
        print('=' * 50)
        print('  CONCLUIDO! index.html atualizado.')
        print('  Rode publicar.bat para enviar ao GitHub.')
        print('=' * 50)
        print()
    except Exception as e:
        print(f'\n[ERRO] {e}')
        import traceback
        traceback.print_exc()


if __name__ == '__main__':
    main()
